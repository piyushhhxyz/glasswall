"""Turn an S3 payload into something the viewer can put side by side."""

from __future__ import annotations

import bz2
import csv
import gzip
import hashlib
import io
import json
import lzma
import re
import tempfile
import zipfile
import zlib
from dataclasses import dataclass, field
from typing import Callable

# Rendered by the browser rather than diffed as text. The redactor paints over
# PII inside images, so seeing the two versions is the whole point for these.
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".tif", ".svg",
                  ".ico", ".heic", ".heif", ".avif"}

# No useful text and no useful rendering: report byte identity and stop.
OPAQUE_SUFFIXES = {
    ".mp3", ".mp4", ".m4a", ".mov", ".avi", ".mkv", ".wav", ".flac", ".ogg", ".webm",
    ".woff", ".woff2", ".ttf", ".otf", ".eot", ".dmg", ".iso", ".exe", ".dll", ".so",
    ".dylib", ".class", ".jar", ".pyc", ".wasm", ".pack", ".idx",
}

PLAIN_TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".log", ".text", ".rst", ".ini", ".cfg", ".conf", ".toml",
    ".env", ".properties", ".gitignore", ".sql", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".java", ".kt", ".go", ".rb", ".php", ".c", ".h", ".cpp", ".cs", ".rs", ".swift",
    ".sh", ".bash", ".zsh", ".yaml", ".yml", ".vtt", ".srt", ".ics", ".ical", ".vcf",
    ".patch", ".diff", ".tsv2", ".graphql", ".proto", ".tf", ".lock",
}

# Wrappers: decompress, then dispatch again on the name underneath.
COMPRESSED: dict[str, Callable[[bytes], bytes]] = {
    ".gz": lambda p: gzip.decompress(p),
    ".bz2": lambda p: bz2.decompress(p),
    ".xz": lambda p: lzma.decompress(p),
    ".lzma": lambda p: lzma.decompress(p),
}

MAX_STRINGS = 4000
MIN_STRING_RUN = 4


@dataclass
class Extracted:
    text: str
    kind: str
    note: str | None = None
    truncated: bool = False
    # Set when the browser should render the bytes instead of diffing text.
    media_type: str | None = None
    # Byte-level facts, always available even when nothing could be parsed.
    digest: str | None = None
    byte_count: int = 0
    extras: dict = field(default_factory=dict)

    def to_json(self) -> dict:
        return {
            "text": self.text,
            "kind": self.kind,
            "note": self.note,
            "truncated": self.truncated,
            "media_type": self.media_type,
            "digest": self.digest,
            "byte_count": self.byte_count,
            **self.extras,
        }


EMPTY = Extracted(text="", kind="empty", note="object is empty or missing")


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #

def suffixes(name: str) -> list[str]:
    """All trailing extensions, lowercased: `part.jsonl.gz` -> ['.jsonl', '.gz']."""
    base = name.rsplit("/", 1)[-1]
    parts = base.lower().split(".")
    return [f".{part}" for part in parts[1:]] if len(parts) > 1 else []


def _suffix(name: str) -> str:
    tail = suffixes(name)
    return tail[-1] if tail else ""


# A BOM is the only safe signal for UTF-16/32: without one, `decode("utf-16")`
# succeeds on almost any even-length input and returns mojibake, silently
# destroying the text a scan depends on.
_BOMS = (
    (b"\xff\xfe\x00\x00", "utf-32"),
    (b"\x00\x00\xfe\xff", "utf-32"),
    (b"\xff\xfe", "utf-16"),
    (b"\xfe\xff", "utf-16"),
    (b"\xef\xbb\xbf", "utf-8-sig"),
)


def _decode(payload: bytes) -> tuple[str, str | None]:
    """UTF-8 where possible, lossy where not."""
    for bom, encoding in _BOMS:
        if payload.startswith(bom):
            try:
                return payload.decode(encoding), f"decoded as {encoding}"
            except (UnicodeDecodeError, UnicodeError):
                break
    try:
        return payload.decode("utf-8"), None
    except UnicodeDecodeError:
        return payload.decode("utf-8", errors="replace"), "contained bytes that are not valid UTF-8"


def _looks_binary(payload: bytes) -> bool:
    """A NUL in the first block beats the extension as a binary signal."""
    if any(payload.startswith(bom) for bom, _ in _BOMS):
        return False
    return b"\x00" in payload[:4096]


def _strings(payload: bytes) -> str:
    """Printable runs out of opaque bytes, one per line."""
    runs = re.findall(rb"[\x20-\x7e]{%d,}" % MIN_STRING_RUN, payload)
    return "\n".join(run.decode("ascii", "replace") for run in runs[:MAX_STRINGS])


def _partial_decompress(payload: bytes, wbits: int) -> tuple[bytes, bool]:
    """Inflate as much as possible, tolerating a stream cut off by the byte cap."""
    obj = zlib.decompressobj(wbits)
    try:
        return obj.decompress(payload), False
    except zlib.error:
        # Feed it block by block and keep whatever came out before the break.
        obj = zlib.decompressobj(wbits)
        out = bytearray()
        for start in range(0, len(payload), 8192):
            try:
                out += obj.decompress(payload[start:start + 8192])
            except zlib.error:
                break
        return bytes(out), True


# --------------------------------------------------------------------------- #
# Structured-text extractors
# --------------------------------------------------------------------------- #

def _json(payload: bytes) -> Extracted:
    """Re-indent so a diff lands on fields rather than one enormous line."""
    text, note = _decode(payload)
    try:
        return Extracted(json.dumps(json.loads(text), indent=2, ensure_ascii=False), "json", note)
    except (json.JSONDecodeError, RecursionError):
        # A capped fetch yields invalid JSON by construction. The raw text is
        # still perfectly scannable, so fall back rather than refusing.
        return Extracted(text, "json-invalid", note)


def _jsonl(payload: bytes) -> Extracted:
    """Expand each record onto its own indented block, dropping a trailing partial."""
    text, note = _decode(payload)
    lines = text.splitlines()
    out: list[str] = []
    for index, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        try:
            out.append(json.dumps(json.loads(line), indent=2, ensure_ascii=False))
        except json.JSONDecodeError:
            if index == len(lines) - 1:
                continue  # partial final record from a capped read
            out.append(line)
    return Extracted("\n".join(out), "jsonl", note)


def _delimited(payload: bytes, delimiter: str, kind: str) -> Extracted:
    """CSV/TSV as aligned columns, so a shifted cell is visible in the diff."""
    text, note = _decode(payload)
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error:
        return Extracted(text, f"{kind}-invalid", note)
    if not rows:
        return Extracted(text, kind, note)
    width = min(max(len(row) for row in rows), 64)
    widths = [0] * width
    for row in rows[:2000]:
        for index, cell in enumerate(row[:width]):
            widths[index] = max(widths[index], min(len(cell), 40))
    lines = [
        " | ".join(cell[:60].ljust(widths[i]) for i, cell in enumerate(row[:width]))
        for row in rows
    ]
    return Extracted("\n".join(lines), kind, note, extras={"rows": len(rows)})


def _xml(payload: bytes, kind: str) -> Extracted:
    """Pretty-print via lxml, which the redactor already depends on."""
    text, note = _decode(payload)
    try:
        from lxml import etree

        parser = etree.XMLParser(recover=True, remove_blank_text=True)
        root = etree.fromstring(payload, parser)
        if root is None:
            return Extracted(text, f"{kind}-invalid", note)
        pretty = etree.tostring(root, pretty_print=True, encoding="unicode")
        return Extracted(pretty, kind, note)
    except Exception:  # noqa: BLE001 - malformed markup is expected, fall back to raw
        return Extracted(text, f"{kind}-invalid", note)


def _docx(payload: bytes) -> Extracted:
    """Flatten a DOCX to paragraph text."""
    try:
        from pii_redactor.docx_io import Document

        with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
            handle.write(payload)
            handle.flush()
            document = Document(handle.name)
            return Extracted("\n".join(p.text for p in document.paragraphs()), "docx")
    except Exception as exc:  # noqa: BLE001 - a corrupt or partial docx is expected
        return _ooxml_fallback(payload, "docx", f"could not read docx ({exc})")


def _ooxml_fallback(payload: bytes, kind: str, note: str) -> Extracted:
    """Pull `<w:t>`-style text out of any OOXML package with lxml alone.

    Covers a missing parser or a variant (`.docm`, `.pptm`) with no new dependency.
    """
    try:
        from lxml import etree

        chunks: list[str] = []
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = [n for n in archive.namelist() if n.endswith(".xml")]
            for name in sorted(names):
                if not re.search(r"(document|slide\d*|sheet\d*|shared|notes)", name):
                    continue
                root = etree.fromstring(archive.read(name), etree.XMLParser(recover=True))
                if root is None:
                    continue
                text = " ".join(t.strip() for t in root.itertext() if t.strip())
                if text:
                    chunks.append(f"--- {name} ---\n{text}")
        if chunks:
            return Extracted("\n\n".join(chunks), f"{kind}-xml", note)
    except Exception:  # noqa: BLE001
        pass
    return Extracted(_strings(payload), f"{kind}-strings", f"{note}; showing extracted strings")


def _pdf(payload: bytes) -> Extracted:
    try:
        import pypdf
    except ImportError:
        return Extracted(
            _strings(payload), "pdf-strings",
            "install pypdf for page text; showing extracted strings",
        )
    try:
        reader = pypdf.PdfReader(io.BytesIO(payload), strict=False)
        pages = []
        for number, page in enumerate(reader.pages, 1):
            body = (page.extract_text() or "").strip()
            pages.append(f"--- page {number} ---\n{body}")
        note = "pdf is encrypted; text may be unavailable" if reader.is_encrypted else None
        return Extracted("\n\n".join(pages), "pdf", note, extras={"pages": len(reader.pages)})
    except Exception as exc:  # noqa: BLE001 - a truncated PDF is expected here
        return Extracted(_strings(payload), "pdf-strings", f"could not parse pdf ({exc})")


def _xlsx(payload: bytes) -> Extracted:
    try:
        import openpyxl
    except ImportError:
        return _ooxml_fallback(payload, "xlsx", "install openpyxl for cell values")
    try:
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        out: list[str] = []
        for sheet in book.worksheets:
            out.append(f"--- sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    out.append(" | ".join("" if c is None else str(c) for c in row))
        book.close()
        return Extracted("\n".join(out), "xlsx")
    except Exception as exc:  # noqa: BLE001
        return _ooxml_fallback(payload, "xlsx", f"could not parse xlsx ({exc})")


def _pptx(payload: bytes) -> Extracted:
    try:
        from pptx import Presentation
    except ImportError:
        return _ooxml_fallback(payload, "pptx", "install python-pptx for slide text")
    try:
        deck = Presentation(io.BytesIO(payload))
        out: list[str] = []
        for number, slide in enumerate(deck.slides, 1):
            out.append(f"--- slide {number} ---")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    out.append(shape.text_frame.text)
        return Extracted("\n".join(out), "pptx", extras={"slides": len(deck.slides)})
    except Exception as exc:  # noqa: BLE001
        return _ooxml_fallback(payload, "pptx", f"could not parse pptx ({exc})")


def _parquet(payload: bytes) -> Extracted:
    try:
        import pyarrow.parquet as pq
    except ImportError:
        return Extracted(
            _strings(payload), "parquet-strings",
            "install pyarrow for column values; showing extracted strings",
        )
    try:
        table = pq.read_table(io.BytesIO(payload))
        header = " | ".join(table.schema.names)
        lines = [header, "-" * len(header)]
        for batch in table.to_batches(max_chunksize=200):
            for row in batch.to_pylist():
                lines.append(" | ".join("" if v is None else str(v) for v in row.values()))
        return Extracted("\n".join(lines), "parquet", extras={"rows": table.num_rows})
    except Exception as exc:  # noqa: BLE001 - a capped read cuts the footer off
        return Extracted(
            _strings(payload), "parquet-strings",
            f"could not parse parquet ({exc}); a capped read removes the footer",
        )


def _email(payload: bytes) -> Extracted:
    """Headers plus decoded text parts. Relevant for the `gmail/` prefix."""
    try:
        from email import policy
        from email.parser import BytesParser

        message = BytesParser(policy=policy.default).parsebytes(payload)
        lines = [
            f"{key}: {value}"
            for key, value in message.items()
            if key.lower() in {
                "from", "to", "cc", "bcc", "subject", "date", "reply-to",
                "message-id", "return-path", "sender", "delivered-to",
            }
        ]
        lines.append("")
        for part in message.walk():
            if part.get_content_maintype() != "text":
                name = part.get_filename()
                if name:
                    lines.append(f"[attachment: {name} ({part.get_content_type()})]")
                continue
            try:
                body = part.get_content()
            except Exception:  # noqa: BLE001
                continue
            if isinstance(body, str) and body.strip():
                lines.append(f"--- {part.get_content_type()} ---")
                lines.append(body)
        return Extracted("\n".join(lines), "email")
    except Exception as exc:  # noqa: BLE001
        text, note = _decode(payload)
        return Extracted(text, "email-raw", f"could not parse message ({exc}); {note or 'raw text'}")


def _zip(payload: bytes) -> Extracted:
    """A manifest, plus the text of small members."""
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos = archive.infolist()
            lines = [f"{'size':>10}  {'crc':>10}  name", "-" * 46]
            for info in infos:
                lines.append(f"{info.file_size:>10}  {info.CRC:>10x}  {info.filename}")
            for info in infos:
                if info.file_size > 200_000 or info.is_dir():
                    continue
                if _suffix(info.filename) not in PLAIN_TEXT_SUFFIXES | {".json", ".jsonl", ".csv", ".xml"}:
                    continue
                inner = extract(info.filename, archive.read(info))
                if inner.text:
                    lines.append(f"\n--- {info.filename} ---\n{inner.text}")
            return Extracted("\n".join(lines), "zip", extras={"members": len(infos)})
    except Exception as exc:  # noqa: BLE001 - a capped read truncates the central directory
        return Extracted(
            _strings(payload), "zip-strings",
            f"could not read archive ({exc}); a capped read removes the index",
        )


def _tar(payload: bytes) -> Extracted:
    """Member manifest, plus the text of small members. `.tar.gz` arrives here
    already decompressed by the wrapper handling above."""
    import tarfile

    try:
        with tarfile.open(fileobj=io.BytesIO(payload), mode="r:") as archive:
            members = archive.getmembers()
            lines = [f"{'size':>10}  name", "-" * 40]
            lines += [f"{m.size:>10}  {m.name}" for m in members]
            for member in members:
                if not member.isfile() or member.size > 200_000:
                    continue
                if _suffix(member.name) not in PLAIN_TEXT_SUFFIXES | {".json", ".jsonl", ".csv", ".xml"}:
                    continue
                handle = archive.extractfile(member)
                if handle is None:
                    continue
                inner = extract(member.name, handle.read())
                if inner.text:
                    lines.append(f"\n--- {member.name} ---\n{inner.text}")
            return Extracted("\n".join(lines), "tar", extras={"members": len(members)})
    except Exception as exc:  # noqa: BLE001 - a capped read truncates the archive
        return Extracted(_strings(payload), "tar-strings", f"could not read archive ({exc})")


def _odf(payload: bytes) -> Extracted:
    """OpenDocument text/spreadsheet/presentation: a zip whose body is content.xml."""
    try:
        from lxml import etree

        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            chunks = []
            for name in ("content.xml", "styles.xml", "meta.xml"):
                if name not in archive.namelist():
                    continue
                root = etree.fromstring(archive.read(name), etree.XMLParser(recover=True))
                if root is None:
                    continue
                # ODF marks row/paragraph ends structurally, so join on the
                # element boundaries rather than flattening to one long line.
                text = "\n".join(
                    line for line in (
                        " ".join(t.strip() for t in element.itertext() if t.strip())
                        for element in root.iter()
                        if etree.QName(element).localname in ("p", "h", "table-row")
                    ) if line
                )
                if text:
                    chunks.append(f"--- {name} ---\n{text}")
        if chunks:
            return Extracted("\n\n".join(chunks), "odf")
    except Exception as exc:  # noqa: BLE001
        return Extracted(_strings(payload), "odf-strings", f"could not parse document ({exc})")
    return Extracted(_strings(payload), "odf-strings", "no readable content part")


# RTF groups that hold metadata rather than document text. Everything inside one
# of these is dropped; `\*` marks any other destination the reader may ignore.
_RTF_SKIP_GROUPS = frozenset({
    "fonttbl", "colortbl", "stylesheet", "listtable", "listoverridetable", "info",
    "pict", "object", "themedata", "colorschememapping", "latentstyles", "datastore",
    "generator", "filetbl", "revtbl", "xmlnstbl", "upr", "rsidtbl", "mmathPr",
})
# Control words that stand for whitespace rather than formatting. `pard` resets
# paragraph properties and in practice opens each new block, so treating it as a
# break is what gives a line-based diff something to align on.
_RTF_BREAKS = frozenset({"par", "pard", "line", "row", "sect", "page"})


def _rtf(payload: bytes) -> Extracted:
    """Strip RTF control words down to the visible text."""
    text, note = _decode(payload)
    out: list[str] = []
    depth, skip_from = 0, None
    index, length = 0, len(text)

    while index < length:
        char = text[index]

        if char == "{":
            depth += 1
            # Look at the group's first control word to decide whether to keep it.
            head = re.match(r"\{\\(\*|[a-zA-Z]+)", text[index:])
            if head and skip_from is None:
                word = head.group(1)
                if word == "*" or word in _RTF_SKIP_GROUPS:
                    skip_from = depth
            index += 1
            continue

        if char == "}":
            if skip_from is not None and depth <= skip_from:
                skip_from = None
            depth -= 1
            index += 1
            continue

        if char == "\\":
            escape = text[index + 1: index + 2]
            if escape in ("\\", "{", "}"):
                if skip_from is None:
                    out.append(escape)
                index += 2
                continue
            # A backslash before a literal newline is RTF's line break. TextEdit
            # and browser-generated RTF use it instead of `\par`, so dropping it
            # collapses the whole document onto one line.
            if escape in ("\r", "\n"):
                if skip_from is None:
                    out.append("\n")
                index += 2
                continue
            if escape == "'":  # \'xx -- a single byte in hex
                hex_pair = text[index + 2: index + 4]
                if skip_from is None and re.fullmatch(r"[0-9a-fA-F]{2}", hex_pair or ""):
                    out.append(bytes([int(hex_pair, 16)]).decode("cp1252", "replace"))
                index += 4
                continue
            word = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?", text[index:])
            if word:
                name, argument = word.group(1), word.group(2)
                if skip_from is None:
                    if name == "u" and argument:  # \uNNNN -- a Unicode code point
                        out.append(chr(int(argument) % 0x10000))
                    elif name in _RTF_BREAKS:
                        out.append("\n")
                    elif name in ("tab", "cell"):
                        out.append("\t")
                index += word.end()
                continue
            index += 2  # a lone backslash before something unexpected
            continue

        if skip_from is None and char not in "\r\n":
            out.append(char)
        index += 1

    body = re.sub(r"[ \t]{2,}", " ", "".join(out))
    body = re.sub(r"\n{3,}", "\n\n", body).strip()
    return Extracted(body, "rtf", note)


def _sqlite(payload: bytes) -> Extracted:
    """Schema and row counts. Enough to see a table gained or lost."""
    import sqlite3

    try:
        with tempfile.NamedTemporaryFile(suffix=".db") as handle:
            handle.write(payload)
            handle.flush()
            connection = sqlite3.connect(f"file:{handle.name}?mode=ro", uri=True)
            rows = connection.execute(
                "SELECT type, name, sql FROM sqlite_master ORDER BY type, name"
            ).fetchall()
            lines: list[str] = []
            for kind, name, sql in rows:
                lines.append(f"--- {kind}: {name} ---")
                if sql:
                    lines.append(sql)
                if kind == "table":
                    try:
                        count = connection.execute(f'SELECT COUNT(*) FROM "{name}"').fetchone()[0]
                        lines.append(f"rows: {count}")
                    except sqlite3.Error:
                        pass
            connection.close()
            return Extracted("\n".join(lines), "sqlite")
    except Exception as exc:  # noqa: BLE001
        return Extracted(_strings(payload), "sqlite-strings", f"could not open database ({exc})")


def _image(name: str, payload: bytes) -> Extracted:
    """Rendered by the browser; OCR'd too when macOS Vision is available."""
    suffix = _suffix(name)
    media = {
        ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
        ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp",
        ".tiff": "image/tiff", ".tif": "image/tiff", ".svg": "image/svg+xml",
        ".ico": "image/x-icon", ".heic": "image/heic", ".heif": "image/heif",
        ".avif": "image/avif",
    }.get(suffix, "application/octet-stream")

    # SVG is markup: diff it as text and render it.
    if suffix == ".svg":
        result = _xml(payload, "svg")
        result.media_type = media
        return result

    text, note, kind = "", None, "image"
    try:
        from pii_redactor import images as ocr_module

        if ocr_module.available():
            boxes = ocr_module.ocr(payload)
            text = "\n".join(box.text for box in boxes)
            kind = "image-ocr"
            note = "text below was read out of the image by OCR"
        else:
            note = "install pyobjc (Quartz, Vision) to OCR image text on macOS"
    except Exception as exc:  # noqa: BLE001 - OCR is best-effort
        note = f"OCR unavailable ({exc})"
    return Extracted(text, kind, note, media_type=media)


# --------------------------------------------------------------------------- #
# Dispatch
# --------------------------------------------------------------------------- #

EXTRACTORS: dict[str, Callable[[bytes], Extracted]] = {
    ".json": _json,
    ".geojson": _json,
    ".jsonl": _jsonl,
    ".ndjson": _jsonl,
    ".csv": lambda p: _delimited(p, ",", "csv"),
    ".tsv": lambda p: _delimited(p, "\t", "tsv"),
    ".xml": lambda p: _xml(p, "xml"),
    ".html": lambda p: _xml(p, "html"),
    ".htm": lambda p: _xml(p, "html"),
    ".xhtml": lambda p: _xml(p, "html"),
    ".plist": lambda p: _xml(p, "xml"),
    ".rels": lambda p: _xml(p, "xml"),
    ".docx": _docx,
    ".docm": _docx,
    ".pdf": _pdf,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".pptx": _pptx,
    ".pptm": _pptx,
    ".parquet": _parquet,
    ".eml": _email,
    ".mbox": _email,
    ".zip": _zip,
    ".apk": _zip,
    ".jar": _zip,
    ".tar": _tar,
    ".odt": _odf,
    ".ods": _odf,
    ".odp": _odf,
    ".odg": _odf,
    ".rtf": _rtf,
    ".sqlite": _sqlite,
    ".sqlite3": _sqlite,
    ".db": _sqlite,
    ".avro": lambda p: Extracted(_strings(p), "avro-strings", "showing extracted strings"),
    ".doc": lambda p: Extracted(_strings(p), "doc-strings",
                                "legacy .doc has no parser here; showing extracted strings"),
    ".xls": lambda p: Extracted(_strings(p), "xls-strings",
                                "legacy .xls has no parser here; showing extracted strings"),
}


def extract(name: str, payload: bytes, truncated: bool = False) -> Extracted:
    """Best-effort text for `payload`, named by `name` so the suffix can route it."""
    if not payload:
        return EMPTY

    result = _extract_inner(name, payload)
    result.truncated = result.truncated or truncated
    result.byte_count = len(payload)
    result.digest = hashlib.sha256(payload).hexdigest()[:16]
    return result


def _extract_inner(name: str, payload: bytes) -> Extracted:
    tail = suffixes(name)
    suffix = tail[-1] if tail else ""

    # Compression wrapper: unwrap, then dispatch on the name underneath.
    if suffix in COMPRESSED:
        inner_name = name[: -len(suffix)]
        partial = False
        try:
            payload = COMPRESSED[suffix](payload)
        except Exception:  # noqa: BLE001 - a capped read cuts the stream mid-member
            if suffix == ".gz":
                payload, partial = _partial_decompress(payload, 16 + zlib.MAX_WBITS)
            else:
                payload = b""
            if not payload:
                return Extracted(
                    "", f"{suffix.lstrip('.')}-error",
                    f"could not decompress; a capped read cuts the {suffix} stream",
                )
        result = _extract_inner(inner_name, payload)
        result.kind = f"{result.kind}+{suffix.lstrip('.')}"
        if partial:
            result.truncated = True
            result.note = (result.note + "; " if result.note else "") + \
                "decompressed only as far as the byte cap allowed"
        return result

    if suffix in IMAGE_SUFFIXES:
        return _image(name, payload)

    if suffix in OPAQUE_SUFFIXES:
        return Extracted(
            "", "opaque",
            f"{suffix.lstrip('.')} content is compared by size and digest only",
            media_type="application/octet-stream",
        )

    extractor = EXTRACTORS.get(suffix)
    if extractor:
        return extractor(payload)

    if suffix in PLAIN_TEXT_SUFFIXES:
        # The extension is not evidence: the export contains `.txt` names holding
        # binary payloads. Content wins.
        if _looks_binary(payload):
            return _sniff(payload)
        text, note = _decode(payload)
        return Extracted(text, suffix.lstrip("."), note)

    # Unknown suffix. Sniff the content rather than giving up on it.
    return _sniff(payload)


def _sniff(payload: bytes) -> Extracted:
    """Identify by content when the name gives nothing away."""
    if payload[:4] == b"%PDF":
        return _pdf(payload)
    if payload[:2] == b"PK":
        # An OOXML or ODF package is a zip with a known part inside it.
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                names = set(archive.namelist())
            if "word/document.xml" in names:
                return _docx(payload)
            if any(n.startswith("ppt/slides/") for n in names):
                return _pptx(payload)
            if "xl/workbook.xml" in names:
                return _xlsx(payload)
            if "content.xml" in names and "META-INF/manifest.xml" in names:
                return _odf(payload)
        except Exception:  # noqa: BLE001
            pass
        return _zip(payload)
    if payload[:5] == b"{\\rtf":
        return _rtf(payload)
    if payload[257:262] == b"ustar":
        return _tar(payload)
    if payload[:6] in (b"SQLite",) or payload[:16].startswith(b"SQLite format 3"):
        return _sqlite(payload)
    if payload[:8] == b"\x89PNG\r\n\x1a\n":
        return _image("x.png", payload)
    if payload[:3] == b"\xff\xd8\xff":
        return _image("x.jpg", payload)
    if payload[:4] == b"PAR1":
        return _parquet(payload)
    if payload[:2] == b"\x1f\x8b":
        return _extract_inner("inner.gz", payload)

    if _looks_binary(payload):
        return Extracted(
            _strings(payload), "binary-strings",
            "unrecognised binary; showing printable strings so it can still be scanned",
        )

    text, note = _decode(payload)
    stripped = text.lstrip()
    if stripped[:1] in "{[":
        return _json(payload)
    if stripped[:1] == "<":
        return _xml(payload, "xml")
    return Extracted(text, "text", note)
