import docx
import pytest


@pytest.fixture
def make_docx(tmp_path):
    """Build a small .docx, optionally splitting paragraphs into multiple runs."""

    def _make(paragraphs, name="sample.docx", split_runs=False):
        document = docx.Document()
        for text in paragraphs:
            paragraph = document.add_paragraph()
            if split_runs:
                for i, token in enumerate(text.split(" ")):
                    paragraph.add_run(token if i == 0 else " " + token)
            else:
                paragraph.add_run(text)
        path = tmp_path / name
        document.save(path)
        return path

    return _make
