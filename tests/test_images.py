"""Image OCR stage.

Text baked into a picture is invisible to every other part of the tool. These tests
generate their own images, so they do not depend on the prospectus, and they skip
cleanly on a machine without Apple Vision.
"""

import io

import pytest
from PIL import Image, ImageDraw, ImageFont

from pii_redactor import images as image_ocr
from pii_redactor.model import PiiType
from pii_redactor.policy import Policy
from pii_redactor.vault import Vault

pytestmark = pytest.mark.skipif(not image_ocr.available(), reason="needs macOS Vision")

FONT_PATHS = ["/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Helvetica.ttc"]


def render(lines, size=(760, 460)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = None
    for path in FONT_PATHS:
        try:
            font = ImageFont.truetype(path, 34)
            break
        except OSError:
            continue
    for i, line in enumerate(lines):
        draw.text((26, 26 + i * 58), line, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_ocr_reads_text_from_an_image():
    text = " ".join(b.text for b in image_ocr.ocr(render(["INCOME TAX DEPARTMENT"])))
    assert "INCOME TAX" in text.upper()


def test_identity_card_is_redacted_wholesale():
    """A name on a card scan is in no gazetteer; the card is sensitive as an object."""
    payload = render(["INCOME TAX DEPARTMENT", "Permanent Account Number", "ABCDE1234F", "RAVI KUMAR"])
    new, findings = image_ocr.redact_image(payload, [], Vault(), Policy())
    assert new is not None and findings
    recovered = " ".join(b.text for b in image_ocr.ocr(new)).upper()
    assert "ABCDE1234F" not in recovered
    assert "RAVI KUMAR" not in recovered


def test_identity_number_keeps_its_shape_rather_than_passing_through():
    """`person()` only rewrites alphabetic tokens and handed a PAN straight back."""
    out = Vault().surrogate(PiiType.ID_DOCUMENT, "ABCDE1234F")
    assert out != "ABCDE1234F"
    assert len(out) == len("ABCDE1234F")
    assert [c.isdigit() for c in out] == [c.isdigit() for c in "ABCDE1234F"]


def test_issuer_boilerplate_is_kept_so_the_card_stays_recognisable():
    payload = render(["GOVERNMENT OF INDIA", "ANIL SHARMA"])
    new, _ = image_ocr.redact_image(payload, [], Vault(), Policy())
    recovered = " ".join(b.text for b in image_ocr.ocr(new)).upper()
    assert "GOVERNMENT OF INDIA" in recovered
    assert "ANIL SHARMA" not in recovered


def test_ordinary_image_without_pii_is_left_alone():
    new, findings = image_ocr.redact_image(render(["Annual Production Capacity"]), [], Vault(), Policy())
    assert new is None and findings == []


def test_logo_is_matched_despite_ocr_noise():
    """Logos OCR badly: "Bluecrest" comes back as "_bluecrest". Word boundaries fail."""
    payload = render(["_bluecrest"])
    new, findings = image_ocr.redact_image(payload, [], Vault(), Policy(),
                                           known_terms={"Bluecrest Capital"})
    assert new is not None and findings


def test_qr_code_is_reencoded_and_still_scannable():
    import qrcode
    buffer = io.BytesIO()
    qrcode.make("https://qrfy.io/tfC6dQ2xWg").convert("RGB").save(buffer, format="PNG")
    new, findings = image_ocr.redact_barcode(buffer.getvalue(), Vault())
    assert findings and findings[0]["original"] == "https://qrfy.io/tfC6dQ2xWg"
    decoded = image_ocr.barcodes(new)
    assert decoded and "qrfy.io" not in decoded[0]


def test_pipeline_rewrites_media_parts(tmp_path):
    import zipfile

    import docx
    from pii_redactor.pipeline import redact

    picture = tmp_path / "card.png"
    picture.write_bytes(render(["INCOME TAX DEPARTMENT", "ABCDE1234F", "RAVI KUMAR"]))
    document = docx.Document()
    document.add_paragraph("See attached.")
    document.add_picture(str(picture))
    source = tmp_path / "in.docx"
    document.save(source)

    result = redact(source, tmp_path / "out.docx")
    assert result.images_rewritten == 1
    with zipfile.ZipFile(tmp_path / "out.docx") as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        recovered = " ".join(b.text for b in image_ocr.ocr(zf.read(media[0]))).upper()
    assert "ABCDE1234F" not in recovered


def test_image_stage_is_skippable(tmp_path):
    import docx
    from pii_redactor.pipeline import redact

    picture = tmp_path / "card.png"
    picture.write_bytes(render(["ABCDE1234F"]))
    document = docx.Document()
    document.add_picture(str(picture))
    source = tmp_path / "in.docx"
    document.save(source)
    assert redact(source, tmp_path / "out.docx", redact_images=False).images_rewritten == 0
