"""Getting text out of the formats the export actually contains."""

from __future__ import annotations

import json

import pytest

from pii_verifier.extract import extract


def test_json_is_reindented_so_a_diff_is_readable():
    """Slack exports are minified; one long line makes the diff say nothing."""
    payload = json.dumps({"a": 1, "b": {"c": 2}}).encode()
    result = extract("x.json", payload)
    assert result.kind == "json"
    assert result.text.count("\n") > 1


def test_truncated_json_still_yields_scannable_text():
    """A capped fetch cuts JSON mid-structure; the text is still worth scanning."""
    payload = b'{"text": "mail a@example.com", "more": [1, 2'
    result = extract("x.json", payload, truncated=True)
    assert result.kind == "json-invalid"
    assert "a@example.com" in result.text
    assert result.truncated


def test_jsonl_records_are_expanded_individually():
    payload = b'{"a": 1}\n{"a": 2}\n'
    result = extract("x.jsonl", payload)
    assert result.kind == "jsonl"
    assert result.text.count('"a"') == 2


def test_jsonl_drops_only_a_trailing_partial_record():
    payload = b'{"a": 1}\n{"b": 2}\n{"c": '
    result = extract("x.jsonl", payload, truncated=True)
    assert '"a"' in result.text and '"b"' in result.text
    assert '"c"' not in result.text


def test_image_is_marked_renderable_rather_than_skipped():
    """The redactor paints over PII in pixels, so images must be viewable."""
    result = extract("logo.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    assert result.kind.startswith("image")
    assert result.media_type == "image/png"


def test_binary_is_detected_by_content_when_the_suffix_lies():
    """The export contains `.txt` names holding binary payloads."""
    result = extract("data.txt", b"ok\x00\x01\x02binary-with-a@example.com inside")
    assert result.kind == "binary-strings"
    # Tier 2: PII in an opaque blob is still PII.
    assert "a@example.com" in result.text


def test_opaque_media_reports_byte_identity_only():
    payload = b"\x00\x00\x00 ftypmp42" + b"\x11" * 50
    result = extract("clip.mp4", payload)
    assert result.kind == "opaque"
    assert result.text == ""
    assert result.digest and result.byte_count == len(payload)


def test_utf16_is_only_decoded_behind_a_bom():
    """Without a BOM, UTF-16 "succeeds" on arbitrary bytes and returns mojibake."""
    latin = extract("x.txt", b"caf\xe9 a@example.com")
    assert "a@example.com" in latin.text

    marked = extract("x.txt", b"\xff\xfe" + "café a@example.com".encode("utf-16-le"))
    assert "a@example.com" in marked.text
    assert marked.note == "decoded as utf-16"


def test_invalid_utf8_is_decoded_lossily_with_a_note():
    result = extract("x.txt", b"caf\xe9 a@example.com")
    assert "a@example.com" in result.text
    assert result.note


def test_unknown_suffix_holding_json_is_still_structured():
    result = extract("blob.dat", b'{"text": "a@example.com"}')
    assert result.kind == "json"


def test_empty_payload_is_reported_as_empty():
    assert extract("x.json", b"").kind == "empty"


def test_docx_is_read_from_bytes(tmp_path):
    """`docx_io.Document` needs a path; extraction must bridge that from bytes."""
    import docx

    document = docx.Document()
    document.add_paragraph("Reach person.one@example.com for access.")
    path = tmp_path / "sample.docx"
    document.save(path)

    result = extract("sample.docx", path.read_bytes())
    assert result.kind == "docx", result.note
    assert "person.one@example.com" in result.text


# --------------------------------------------------------------------------- #
# The wider format range
# --------------------------------------------------------------------------- #

PII = "person.one@example.com"


def test_csv_is_aligned_into_columns():
    result = extract("x.csv", b"name,email\nAtmesh,person.one@example.com\n")
    assert result.kind == "csv"
    assert PII in result.text
    assert "|" in result.text


def test_tsv_uses_the_tab_delimiter():
    assert extract("x.tsv", b"a\tb\n1\t2\n").kind == "tsv"


def test_xml_is_pretty_printed():
    result = extract("x.xml", b"<r><a>person.one@example.com</a><b/></r>")
    assert result.kind == "xml"
    assert PII in result.text and result.text.count("\n") >= 1


def test_html_is_handled_as_markup():
    assert extract("x.html", b"<html><body><p>hi</p></body></html>").kind == "html"


def test_gzip_is_unwrapped_and_dispatched_on_the_inner_name():
    import gzip as gz

    payload = gz.compress(json.dumps({"email": PII}).encode())
    result = extract("part.json.gz", payload)
    assert result.kind == "json+gz"
    assert PII in result.text


def test_bzip2_and_xz_wrappers_work():
    import bz2
    import lzma

    assert extract("a.txt.bz2", bz2.compress(b"hello " + PII.encode())).kind == "txt+bz2"
    assert extract("a.txt.xz", lzma.compress(b"hello " + PII.encode())).kind == "txt+xz"


def test_zip_lists_members_and_reads_small_text_ones():
    import io as _io
    import zipfile

    buffer = _io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", f"contact {PII}")
        archive.writestr("big.bin", b"\x00" * 10)
    result = extract("bundle.zip", buffer.getvalue())
    assert result.kind == "zip"
    assert "notes.txt" in result.text and "big.bin" in result.text
    assert PII in result.text


def test_tar_lists_members():
    import io as _io
    import tarfile

    buffer = _io.BytesIO()
    with tarfile.open(fileobj=buffer, mode="w") as archive:
        data = f"contact {PII}".encode()
        info = tarfile.TarInfo("notes.txt")
        info.size = len(data)
        archive.addfile(info, _io.BytesIO(data))
    result = extract("bundle.tar", buffer.getvalue())
    assert result.kind == "tar"
    assert PII in result.text


def test_tar_gz_unwraps_then_reads_the_archive():
    import gzip as gz
    import io as _io
    import tarfile

    buffer = _io.BytesIO()
    with tarfile.open(fileobj=buffer, mode="w") as archive:
        data = b"hello"
        info = tarfile.TarInfo("a.txt")
        info.size = len(data)
        archive.addfile(info, _io.BytesIO(data))
    result = extract("bundle.tar.gz", gz.compress(buffer.getvalue()))
    assert result.kind == "tar+gz"
    assert "a.txt" in result.text


def test_eml_yields_headers_and_body():
    raw = (
        b"From: person.one@example.com\r\nTo: person.two@example.com\r\n"
        b"Subject: Access\r\nDate: Mon, 1 Jan 2024 00:00:00 +0000\r\n"
        b"Content-Type: text/plain\r\n\r\nCall +91 98765 43210.\r\n"
    )
    result = extract("mail.eml", raw)
    assert result.kind == "email"
    assert "Subject: Access" in result.text
    assert PII in result.text and "98765" in result.text


def test_rtf_control_words_are_stripped():
    raw = rb"{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}\f0\fs24 Contact person.one@example.com\par}"
    result = extract("doc.rtf", raw)
    assert result.kind == "rtf"
    assert PII in result.text
    assert "rtf1" not in result.text and "fonttbl" not in result.text


def test_rtf_outer_group_is_not_treated_as_metadata():
    """The document itself opens with `{\\rtf1`; a regex broad enough to drop
    `{\\fonttbl...}` also drops the whole file."""
    raw = rb"{\rtf1\ansi{\fonttbl{\f0 Arial;}}{\*\expandedcolortbl;;}\f0 Visible text\par}"
    assert "Visible text" in extract("doc.rtf", raw).text


def test_rtf_escaped_newline_is_a_line_break():
    """TextEdit and browser RTF use `\\` + newline rather than `\\par`."""
    raw = b"{\\rtf1\\ansi \\f0 First line\\\nSecond line\\\n}"
    lines = [line for line in extract("doc.rtf", raw).text.splitlines() if line.strip()]
    assert lines == ["First line", "Second line"]


def test_rtf_hex_and_unicode_escapes_are_decoded():
    raw = rb"{\rtf1\ansi \f0 caf\'e9 \u8364?}"
    text = extract("doc.rtf", raw).text
    assert "café" in text and "€" in text


def test_xlsx_rows_are_read(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    book.active.append(["name", "email"])
    book.active.append(["Atmesh", PII])
    path = tmp_path / "s.xlsx"
    book.save(path)
    result = extract("s.xlsx", path.read_bytes())
    assert result.kind == "xlsx", result.note
    assert PII in result.text


def test_pdf_pages_are_read():
    pypdf = pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = __import__("io").BytesIO()
    writer.write(buffer)
    result = extract("d.pdf", buffer.getvalue())
    # A blank page yields no text, but it must parse as a PDF, not fall to strings.
    assert result.kind == "pdf", result.note
    assert result.extras.get("pages") == 1


def test_parquet_rows_are_read():
    pa = pytest.importorskip("pyarrow")
    import pyarrow.parquet as pq

    buffer = __import__("io").BytesIO()
    pq.write_table(pa.table({"email": [PII]}), buffer)
    result = extract("t.parquet", buffer.getvalue())
    assert result.kind == "parquet", result.note
    assert PII in result.text


def test_sqlite_schema_is_dumped(tmp_path):
    import sqlite3

    path = tmp_path / "d.sqlite"
    connection = sqlite3.connect(path)
    connection.execute("CREATE TABLE people (email TEXT)")
    connection.execute("INSERT INTO people VALUES (?)", (PII,))
    connection.commit()
    connection.close()
    result = extract("d.sqlite", path.read_bytes())
    assert result.kind == "sqlite", result.note
    assert "people" in result.text and "rows: 1" in result.text


# --------------------------------------------------------------------------- #
# Content sniffing, for names that give nothing away
# --------------------------------------------------------------------------- #

def test_extensionless_json_is_still_structured():
    """239 objects in the export have no extension at all."""
    result = extract("no-extension-here", json.dumps({"email": PII}).encode())
    assert result.kind == "json"
    assert PII in result.text


def test_a_company_name_is_not_mistaken_for_an_extension():
    """Redacted names like `report.some logistics corp` end in a fake suffix."""
    result = extract("report.some logistics corp", json.dumps({"a": PII}).encode())
    assert result.kind == "json"


def test_docx_is_sniffed_from_an_ooxml_package(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph(f"Reach {PII}")
    path = tmp_path / "s.docx"
    document.save(path)
    result = extract("mystery-file", path.read_bytes())
    assert result.kind == "docx"
    assert PII in result.text


def test_pdf_is_sniffed_from_its_magic_bytes():
    result = extract("mystery", b"%PDF-1.4\n" + b"\x00" * 20)
    assert result.kind.startswith("pdf")


def test_png_is_sniffed_from_its_magic_bytes():
    result = extract("mystery", b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    assert result.kind.startswith("image")


def test_every_result_carries_byte_facts():
    """Byte identity is the tier-3 fallback and must always be present."""
    for name, payload in [("a.json", b'{"a":1}'), ("a.mp4", b"\x00ftyp"), ("a", b"plain")]:
        result = extract(name, payload)
        assert result.digest and result.byte_count == len(payload)
