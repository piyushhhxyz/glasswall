"""One test per defect found in review. Each fails on the code that shipped before it.

These are the expensive bugs -- the ones a sample-based evaluation could not see,
because they live in parts of the document a reader never visits or in inputs the
sample happened not to contain.
"""

import re
import zipfile

import docx
import pytest
from lxml import etree

from pii_redactor.docx_io import Document, w
from pii_redactor.model import PiiType, Span
from pii_redactor.pipeline import redact
from pii_redactor.policy import Policy, _hostname, _is_institution
from pii_redactor.vault import Vault, _brand_key


@pytest.fixture
def field_code_docx(tmp_path):
    """A .docx whose hyperlink target lives in a field code, as Word writes it.

        <w:instrText> HYPERLINK "mailto:info@acme.com" </w:instrText>

    The address is live but invisible to any reader that only walks <w:t>.
    """
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Email: ")
    run = paragraph.add_run()
    instr = etree.SubElement(run._r, w("instrText"))
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = ' HYPERLINK "mailto:info@acme.com"'
    paragraph.add_run("info@acme.com")
    path = tmp_path / "fields.docx"
    document.save(path)
    return path


def raw_xml(path):
    with zipfile.ZipFile(path) as zf:
        return "".join(zf.read(n).decode("utf8", "replace")
                       for n in zf.namelist() if n.endswith((".xml", ".rels")))


# --- BLOCKING: hyperlink field codes carried live addresses past the redactor ----- #

def test_field_code_target_is_redacted(field_code_docx, tmp_path):
    redact(field_code_docx, tmp_path / "out.docx")
    assert "info@acme.com" not in raw_xml(tmp_path / "out.docx")


def test_field_code_survives_as_a_valid_hyperlink(field_code_docx, tmp_path):
    redact(field_code_docx, tmp_path / "out.docx")
    targets = re.findall(r'HYPERLINK[^<"]*"([^"]+)"', raw_xml(tmp_path / "out.docx"))
    assert targets and targets[0].startswith("mailto:") and "@" in targets[0]


def test_flattener_exposes_field_code_text(field_code_docx):
    assert "HYPERLINK" in Document(field_code_docx).paragraphs()[0].text


def test_two_urls_in_one_field_code_are_both_replaced(tmp_path):
    """`HYPERLINK "http://www.acme.com/"www.acme.com` packs two URLs with no space
    between them; a path pattern that swallowed the quote redacted only the first."""
    document = docx.Document()
    document.add_paragraph().add_run('HYPERLINK "http://www.acme.com/"www.acme.com')
    path = tmp_path / "two.docx"
    document.save(path)
    redact(path, tmp_path / "out.docx")
    assert "acme.com" not in raw_xml(tmp_path / "out.docx")


# --- Brand coherence: a company and its domain must agree ------------------------ #

def test_company_name_and_its_domain_share_a_brand():
    vault = Vault(org_keys=[_brand_key("Acme Industries")])
    brand = vault.org("Acme Industries Limited").split()[0].lower()
    assert brand in vault.url("www.acmeindustries.com")
    assert brand in vault.email("info@acmeindustries.com")


def test_distinct_hosts_never_share_a_surrogate_domain():
    """Two banks collapsed onto one fake host when sub-domains were mishandled."""
    vault = Vault()
    assert vault.domain("sbi.co.in") != vault.domain("federalbank.co.in")
    assert vault.domain("northbank.com") != vault.domain("lawfirm.com")


def test_domain_typos_still_fold_together():
    vault = Vault()
    assert vault.domain("acmeindustries.com") == vault.domain("acmeindsutries.com")


def test_url_query_string_domain_is_rewritten():
    """A real host can ride along in a query parameter: ...?domain=bluecrest.com"""
    out = Vault().url("https://x.com/s/abc?domain=bluecrest.com")
    assert "bluecrest.com" not in out


# --- Institution allowlist must not depend on how the URL is dressed ------------- #

@pytest.mark.parametrize("url", [
    "www.bseindia.com", "http://www.bseindia.com/", "https://www.sebi.gov.in/x?y=1",
    "https://siportal.sebi.gov.in/", "sebi.gov.in",
])
def test_institution_urls_are_kept_whatever_the_scheme(url):
    assert not Policy().accepts(Span(0, len(url), PiiType.URL, url, "t"))


def test_non_institution_url_is_still_redacted():
    url = "https://www.acmeindustries.com/"
    assert Policy().accepts(Span(0, len(url), PiiType.URL, url, "t"))


def test_hostname_normalisation():
    assert _hostname("https://WWW.Sebi.Gov.In/a/b?c=d") == "sebi.gov.in"
    assert _is_institution("siportal.sebi.gov.in")
    assert not _is_institution("notsebi.gov.in.evil.com")


# --- Surrogate fidelity ---------------------------------------------------------- #

def test_sebi_registration_surrogate_keeps_its_format():
    out = Vault().surrogate(PiiType.SEBI_REGN, "INM000013004")
    assert re.fullmatch(r"IN[A-Z]\d{9}", out), out


def test_address_tail_does_not_swallow_the_next_field_label(tmp_path):
    """The address grew right through "Telephone:" and ate the phone detector's cue."""
    text = "163, 5th Floor, Backbay Reclamation, Mumbai – 400020 Telephone: 022-68052182"
    document = docx.Document()
    document.add_paragraph().add_run(text)
    path = tmp_path / "addr.docx"
    document.save(path)
    result = redact(path, tmp_path / "out.docx")
    out = Document(result.output).paragraphs()[0].text
    assert "Telephone:" in out
    assert "022-68052182" not in out
    assert {PiiType.ADDRESS, PiiType.PHONE} <= {d.span.type for d in result.detections}

